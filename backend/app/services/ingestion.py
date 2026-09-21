from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def _normalize_metadata(docs: list[Document], path: Path) -> list[Document]:
    """Store just the filename as source metadata — never the full filesystem path.

    Prevents leaking local paths (like C:\\Users\\...\\file.md) as citation labels
    when the app is served from a different filesystem later.
    """
    filename = path.name
    for d in docs:
        d.metadata["source"] = filename
        d.metadata["filename"] = filename
    return docs


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        docs = TextLoader(str(path), encoding="utf-8").load()
    elif suffix == ".docx":
        docx = DocxDocument(str(path))
        text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
        docs = [Document(page_content=text, metadata={})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return _normalize_metadata(docs, path)



def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))
